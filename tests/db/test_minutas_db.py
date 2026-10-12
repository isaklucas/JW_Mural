"""Integração do Modo Minutas com o banco (mongomock): minutas, arquivo modelo,
documentos gerados e o histórico dos publicadores."""
import importlib

import pytest
from docx import Document

from process.minutas import padrao, periodo


@pytest.fixture
def servico(db_ops, monkeypatch):
    """MinutaService apontando para o db_ops do teste (banco em memória)."""
    mod = importlib.import_module("services.minuta_service")  # `services.minuta_service` (atributo) é a instância
    monkeypatch.setattr(mod, "db_ops", db_ops)
    monkeypatch.setattr(mod, "PASTA_SAIDA", None)  # definido por teste via tmp_path
    return mod


def test_configuracao_roundtrip(db_ops):
    assert db_ops.obter_configuracao("congregacao", "") == ""
    db_ops.salvar_configuracao("congregacao", "Praias Rasas")
    assert db_ops.obter_configuracao("congregacao") == "Praias Rasas"


def test_salvar_minuta_incrementa_versao_e_guarda_anterior(db_ops):
    m = padrao.minuta_padrao("fds_oradores")
    assert db_ops.salvar_minuta(m)["versao"] == 1
    m["nome"] = "Editada"
    assert db_ops.salvar_minuta(m)["versao"] == 2
    assert db_ops.buscar_minuta("fds_oradores")["nome"] == "Editada"
    assert db_ops.db["minutas_versoes"].count_documents({"chave": "fds_oradores"}) == 1


def test_modelo_e_documento_gerado_guardam_bytes(db_ops):
    db_ops.salvar_minuta_modelo("fds_oradores", b"\x00\x01docx", "oficial.docx")
    assert db_ops.buscar_minuta_modelo("fds_oradores") == b"\x00\x01docx"
    assert db_ops.info_minuta_modelo("fds_oradores")["nome"] == "oficial.docx"

    db_ops.salvar_documento_gerado(2026, 11, "fds_oradores", "a.docx", b"conteudo", 1)
    db_ops.salvar_documento_gerado(2026, 11, "fds_oradores", "a.docx", b"novo", 2)  # substitui
    doc = db_ops.buscar_documento_gerado(2026, 11, "fds_oradores")
    assert doc["arquivo"] == b"novo" and doc["minuta_versao"] == 2
    lista = db_ops.listar_documentos_gerados(2026, 11)
    assert len(lista) == 1 and "arquivo" not in lista[0]


def test_garantir_padroes_semeia_minutas_e_modelos(servico, db_ops):
    servico.MinutaService().garantir_padroes()
    for chave in padrao.MINUTAS_PADRAO:
        assert db_ops.buscar_minuta(chave)["versao"] == 1
        assert db_ops.buscar_minuta_modelo(chave)[:2] == b"PK"  # zip/docx


def test_minuta_em_formato_antigo_e_substituida(servico, db_ops):
    db_ops.salvar_minuta({"chave": "fds_oradores", "tabelas": [{"colunas": []}]})  # formato 1
    servico.MinutaService().garantir_padroes()
    assert db_ops.buscar_minuta("fds_oradores")["formato"] == padrao.FORMATO


def test_trocar_modelo_valida_o_arquivo(servico, db_ops):
    svc = servico.MinutaService()
    assert not svc.trocar_modelo("fds_oradores", b"lixo")["success"]
    modelo = svc.obter_modelo("fds_sentinela")
    assert svc.trocar_modelo("fds_oradores", modelo, "outro.docx")["success"]
    assert svc.info_modelo("fds_oradores")["nome"] == "outro.docx"


def test_salvar_minuta_confere_tabelas_do_modelo(servico):
    svc = servico.MinutaService()
    m = padrao.minuta_padrao("fds_sentinela")
    m["tabelas"][0]["modelo"] = 9
    r = svc.salvar_minuta(m)
    assert not r["success"] and "não tem a tabela" in r["message"]

def test_minuta_invalida_no_banco_cai_na_padrao(servico, db_ops):
    estragada = padrao.minuta_padrao("fds_sentinela")
    estragada["tabelas"] = []
    db_ops.db["minutas"].insert_one(estragada)
    m = servico.MinutaService().obter_minuta("fds_sentinela")
    assert m["tabelas"] == padrao.MINUTA_SENTINELA["tabelas"]


def _meses_preenchidos():
    meses = periodo.montar_periodo(2026, 12, 2)  # vira o ano
    for m in meses:
        for s in m["semanas"]:
            s.update(presidente="Carlos Souza", leitor_sentinela="Davi Lima",
                     orador="Visitante", tema_discurso="Tema")
    return meses


def test_gerar_final_semana_salva_reuniao_documentos_e_historico(servico, db_ops, seed, tmp_path, monkeypatch):
    seed(["Carlos Souza", "Davi Lima"])
    monkeypatch.setattr(servico, "PASTA_SAIDA", str(tmp_path))
    svc = servico.MinutaService()
    svc.salvar_congregacao("Praias Rasas")

    caminhos = svc.gerar_final_semana(_meses_preenchidos(), "Dirigente", "FDS teste")

    assert len(caminhos) == 2
    for c in caminhos:
        assert Document(c).tables  # arquivo válido em disco
    reuniao = db_ops.buscar_reuniao_final_semana(2026, 12)
    assert reuniao["modo"] == "minuta" and reuniao["periodo"] == "Dezembro 2026–Janeiro 2027"
    assert len(db_ops.listar_documentos_gerados(2026, 12)) == 2

    # Ano certo no histórico mesmo com o bimestre virando o ano.
    hist = [h["data"] for h in db_ops.buscar_historico_publicador("Carlos Souza")]
    assert "2 de janeiro de 2027" in hist and "5 de dezembro de 2026" in hist

    # Reabrir do banco reproduz o mesmo arquivo.
    exportado = svc.exportar_documento(2026, 12, "fds_oradores")
    assert Document(exportado).tables[0].rows[0].cells[0].text.startswith("Dezembro")

    # A estrutura por mês é reconstruída a partir do que foi salvo (edição).
    meses = svc.meses_da_reuniao(reuniao)
    assert [(m["ano"], m["mes"], len(m["semanas"])) for m in meses] == [(2026, 12, 4), (2027, 1, 5)]


def test_regerar_nao_duplica_historico_e_excluir_limpa_documentos(servico, db_ops, seed, tmp_path, monkeypatch):
    seed(["Carlos Souza", "Davi Lima"])
    monkeypatch.setattr(servico, "PASTA_SAIDA", str(tmp_path))
    svc = servico.MinutaService()

    svc.gerar_final_semana(_meses_preenchidos(), "Dirigente", "FDS")
    svc.gerar_final_semana(_meses_preenchidos(), "Dirigente", "FDS")

    hist = db_ops.buscar_historico_publicador("Carlos Souza")
    assert len(hist) == 9  # 4 sábados de dezembro + 5 de janeiro, sem duplicar

    assert db_ops.excluir_reuniao_final_semana(2026, 12)["success"]
    assert db_ops.buscar_historico_publicador("Carlos Souza") == []
    assert db_ops.listar_documentos_gerados(2026, 12) == []
