"""Modo Minutas (SDD 18) — partes puras: período, validação, renderização a partir do
modelo oficial e extração de layout para a pré-visualização."""
import datetime
import io
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import qn

from process.minutas import padrao, periodo, preview, render

RAIZ = Path(__file__).resolve().parents[1]


def _modelo(chave):
    return (RAIZ / "assets" / "minutas" / padrao.ARQUIVOS_MODELO[chave]).read_bytes()


# ----------------------------------------------------------------- período

def test_meses_do_periodo_vira_o_ano():
    assert periodo.meses_do_periodo(2026, 12, 2) == [(2026, 12), (2027, 1)]
    assert periodo.meses_do_periodo(2026, 11, 2) == [(2026, 11), (2026, 12)]


def test_sabados_do_mes():
    assert [d.day for d in periodo.dias_da_reuniao(2026, 11)] == [7, 14, 21, 28]
    assert [d.day for d in periodo.dias_da_reuniao(2026, 10)] == [3, 10, 17, 24, 31]


def test_url_wol_usa_semana_iso():
    # wol.jw.org/.../2026/45 = 2–8 de novembro de 2026 (conferido no site)
    assert periodo.url_semana_wol(datetime.date(2026, 11, 7)).endswith("/lp-t/2026/45")
    # sábado 2/1/2027 ainda pertence à semana ISO 53 de 2026
    assert periodo.url_semana_wol(datetime.date(2027, 1, 2)).endswith("/lp-t/2026/53")


def test_montar_periodo_gera_semanas_com_rotulo_de_historico():
    meses = periodo.montar_periodo(2026, 12, 2)
    assert [(m['ano'], m['mes']) for m in meses] == [(2026, 12), (2027, 1)]
    primeira_jan = meses[1]['semanas'][0]
    assert primeira_jan['data'] == "2027-01-02"
    assert primeira_jan['data_historico'] == "2 de janeiro de 2027"


def test_nome_periodo():
    assert periodo.nome_periodo([(2026, 11), (2026, 12)]) == "Novembro–Dezembro 2026"
    assert periodo.nome_periodo([(2026, 12), (2027, 1)]) == "Dezembro 2026–Janeiro 2027"


# --------------------------------------------------------------- validação

def test_minutas_padrao_sao_validas_contra_o_modelo():
    for chave in padrao.MINUTAS_PADRAO:
        desc = render.descrever_modelo(_modelo(chave))
        assert padrao.validar_minuta(padrao.minuta_padrao(chave), desc) == []


def test_minuta_padrao_devolve_copia():
    m = padrao.minuta_padrao('fds_oradores')
    m['tabelas'].clear()
    assert padrao.MINUTAS_PADRAO['fds_oradores']['tabelas']


@pytest.mark.parametrize("estragar, trecho", [
    (lambda m: m.update(tabelas=[]), "ao menos um bloco"),
    (lambda m: m['tabelas'][0]['campos'].__setitem__(0, 'inexistente'), "campo desconhecido"),
    (lambda m: m['tabelas'][0].update(modelo=99), "não tem a tabela"),
    (lambda m: m.update(formato=1), "formato antigo"),
])
def test_validacao_aponta_erros(estragar, trecho):
    m = padrao.minuta_padrao('fds_oradores')
    estragar(m)
    desc = render.descrever_modelo(_modelo('fds_oradores'))
    assert any(trecho in e for e in padrao.validar_minuta(m, desc))


# ------------------------------------------------------------- substituições

def test_variavel_vazia_nao_deixa_congregacao_pendurada():
    titulo = "{mes} – Recebendo - Congregação {congregacao_recebendo}"
    assert render.substituir_variaveis(titulo, {'mes': 'Novembro', 'congregacao_recebendo': ''}) == "Novembro – Recebendo"
    assert render.substituir_variaveis(titulo, {'mes': 'Novembro', 'congregacao_recebendo': 'Jacaré'}) \
        == "Novembro – Recebendo - Congregação Jacaré"


def test_valor_dirigente_leitor():
    dados = {'dirigente': 'Ana'}
    assert render.valor_campo('dirigente_leitor', {'leitor_sentinela': 'Bia'}, dados) == "Ana / Bia"
    assert render.valor_campo('dirigente_leitor', {'leitor_sentinela': ''}, dados) == "Ana"
    assert render.valor_campo('data', {'data': '2026-11-07'}, dados) == "07/11/2026"
    assert render.valor_campo('', {}, dados) == ""


def _paragrafo(xml_runs):
    return parse_xml(
        '<w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">' + xml_runs + '</w:p>')


def test_troca_de_texto_preserva_formatacao_por_trecho():
    """"Setembro" 14pt + " – Saindo - Congregação Jacaré" azul 12pt: o mês novo
    herda a formatação do mês antigo e o resto, a do resto."""
    p = _paragrafo(
        '<w:r><w:rPr><w:sz w:val="28"/></w:rPr><w:t>Setembro</w:t></w:r>'
        '<w:r><w:rPr><w:color w:val="3333FF"/><w:sz w:val="24"/></w:rPr>'
        '<w:t xml:space="preserve"> – Saindo - Congregação Jacaré</w:t></w:r>')
    render.definir_texto_paragrafo(p, "Novembro – Saindo - Congregação Vizinha")
    runs = [(r.find(qn('w:rPr')).find(qn('w:sz')).get(qn('w:val')), render._texto_run(r))
            for r in p.iter(qn('w:r'))]
    assert "".join(t for _, t in runs) == "Novembro – Saindo - Congregação Vizinha"
    assert runs[0] == ("28", "Novembro")
    assert all(sz == "24" for sz, _ in runs[1:])


def test_troca_de_texto_nao_mexe_em_imagem_e_tira_hyperlink():
    p = _paragrafo(
        '<w:r><w:drawing/></w:r>'
        '<w:hyperlink><w:r><w:rPr><w:b/></w:rPr><w:t>antigo</w:t></w:r></w:hyperlink>')
    render.definir_texto_paragrafo(p, "novo")
    assert p.find('.//' + qn('w:drawing')) is not None
    assert p.find(qn('w:hyperlink')) is None
    assert render.texto_paragrafo(p) == "novo"


# -------------------------------------------------------------- renderização

def _dados():
    meses = periodo.montar_periodo(2026, 11, 2)
    for m in meses:
        m['congregacao_recebendo'] = 'Praias Rasas'
        for i, s in enumerate(m['semanas'], start=1):
            s.update(presidente=f"Pres {i}", orador=f"Orador {i}", tema_discurso=f"Tema {i}",
                     leitor_sentinela=f"Leitor {i}", titulo_estudo=f"Estudo {i}")
    meses[0]['congregacao_saindo'] = 'Jacaré'
    meses[0]['semanas'][0].update(orador_saindo="Sai 1", tema_saindo="Tema sai 1")
    return {'congregacao': 'Praias Rasas', 'dirigente': 'Dirigente X', 'meses': meses}


def _linhas(tabela):
    return [[" ".join(c.text.split()) for c in r.cells] for r in tabela.rows]


def test_render_oradores_clona_tabelas_do_modelo_por_mes():
    conteudo = render.renderizar(padrao.minuta_padrao('fds_oradores'), _modelo('fds_oradores'), _dados())
    doc = Document(io.BytesIO(conteudo))
    tabelas = [_linhas(t) for t in doc.tables]

    # Novembro: recebendo + saindo; Dezembro: só recebendo (saindo vazio é omitido)
    assert len(tabelas) == 3
    assert tabelas[0][0][0] == "Novembro – Recebendo - Congregação Praias Rasas"
    assert tabelas[0][1] == ["SÁBADO", "PRESIDENTE", "ORADOR", "TEMA: DISCURSO PUBLICO"]
    assert tabelas[0][2] == ["07/11/2026", "Pres 1", "Orador 1", "Tema 1"]
    assert len(tabelas[0]) == 2 + 4  # título + cabeçalho + 4 sábados (sobras do modelo saem)
    assert tabelas[1][0][0] == "Novembro – Saindo - Congregação Jacaré"
    assert tabelas[1][2] == ["07/11/2026", "Sai 1", "Tema sai 1"]
    assert tabelas[2][0][0].startswith("Dezembro")

    # Estilo e cabeçalho (com a imagem) vêm do modelo oficial.
    assert doc.tables[0].style.name == "Grid Table 6 Colorful Accent 5"
    assert doc.tables[1].style.name == "Grid Table 4 Accent 5"
    assert "Reunião de Final de Semana Oradores" in " ".join(p.text for p in doc.paragraphs)
    assert doc.element.body.find('.//' + qn('w:drawing')) is not None
    assert any(r.reltype.endswith('/image') for r in doc.part.rels.values())


def test_render_mantem_formatacao_das_celulas_do_modelo():
    conteudo = render.renderizar(padrao.minuta_padrao('fds_oradores'), _modelo('fds_oradores'), _dados())
    doc = Document(io.BytesIO(conteudo))
    titulo = doc.tables[0].rows[0].cells[0]
    run = next(r for p in titulo.paragraphs for r in p.runs if r.text.strip())
    assert run.font.name == "Verdana" and run.font.size.pt == 14
    assert str(run.font.color.rgb) == "3333FF"
    presidente = next(r for p in doc.tables[0].rows[2].cells[1].paragraphs for r in p.runs)
    assert presidente.bold and presidente.italic and presidente.font.size.pt == 9


def test_render_sentinela_junta_dirigente_e_leitor_e_mantem_titulo():
    conteudo = render.renderizar(padrao.minuta_padrao('fds_sentinela'), _modelo('fds_sentinela'), _dados())
    tabelas = [_linhas(t) for t in Document(io.BytesIO(conteudo)).tables]
    assert len(tabelas) == 2
    assert tabelas[0][0][0].startswith("Estudo de A Sentinela")  # título do modelo mantido
    assert tabelas[0][2] == ["07/11/2026", "Dirigente X / Leitor 1", "Estudo 1"]


def test_render_cabecalho_com_variavel():
    m = padrao.minuta_padrao('fds_oradores')
    m['cabecalho'] = ["Reunião de Final de Semana Oradores Congregação {congregacao}"]
    dados = _dados()
    dados['congregacao'] = 'Jacaré'
    doc = Document(io.BytesIO(render.renderizar(m, _modelo('fds_oradores'), dados)))
    assert any(p.text.endswith("Congregação Jacaré") for p in doc.paragraphs)
    assert doc.element.body.find('.//' + qn('w:drawing')) is not None  # imagem continua


def test_render_bloco_com_tabela_inexistente_e_ignorado():
    m = padrao.minuta_padrao('fds_sentinela')
    m['tabelas'].append({**m['tabelas'][0], 'modelo': 42})
    doc = Document(io.BytesIO(render.renderizar(m, _modelo('fds_sentinela'), _dados())))
    assert len(doc.tables) == 2


def test_validar_modelo():
    assert render.validar_modelo(_modelo('fds_oradores')) is None
    assert "válido" in render.validar_modelo(b"nao sou um docx")
    vazio = io.BytesIO()
    Document().save(vazio)
    assert "tabela" in render.validar_modelo(vazio.getvalue())


def test_preview_extrai_imagem_cores_e_tabelas():
    conteudo = render.renderizar(padrao.minuta_padrao('fds_oradores'), _modelo('fds_oradores'), _dados())
    layout = preview.montar_layout(conteudo)
    assert any(b['tipo'] == 'p' and b['imagens'] for b in layout['blocos'])
    tabela = next(b for b in layout['blocos'] if b['tipo'] == 'tabela')
    titulo = tabela['linhas'][0]['celulas'][0]
    assert titulo['span'] == 4
    runs = [r for p in titulo['paragrafos'] for r in p['runs']]
    assert runs[0]['texto'].startswith("Novembro") and runs[0]['cor'] == "3333FF"
    assert titulo['fundo']  # cor vinda do estilo de tabela (tema)
    assert len(tabela['cor_borda']) == 6
